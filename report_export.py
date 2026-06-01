"""Export an interview report to downloadable PDF and DOCX bytes."""

import io
import re


def _safe_filename(candidate):
    base = re.sub(r"[^A-Za-z0-9]+", "_", candidate or "candidate").strip("_")
    return f"Interview_Report_{base or 'candidate'}"


def build_pdf(report):
    """Return the report as PDF bytes using reportlab."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, HRFlowable)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        topMargin=18 * mm, bottomMargin=18 * mm,
        leftMargin=18 * mm, rightMargin=18 * mm,
        title=f"Interview Report - {report['candidate']}",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("H1c", parent=styles["Heading1"],
                              textColor=colors.HexColor("#1E88E5")))
    styles.add(ParagraphStyle("H2c", parent=styles["Heading2"],
                              textColor=colors.HexColor("#1565C0"), spaceBefore=10))
    styles.add(ParagraphStyle("Body", parent=styles["BodyText"],
                              fontSize=10, leading=14))
    styles.add(ParagraphStyle("Meta", parent=styles["BodyText"],
                              fontSize=9, textColor=colors.HexColor("#555555")))

    def esc(text):
        return (str(text).replace("&", "&amp;")
                .replace("<", "&lt;").replace(">", "&gt;"))

    flow = []
    flow.append(Paragraph(f"Interview Report — {esc(report['candidate'])}", styles["H1c"]))
    flow.append(Paragraph(
        f"Date: {esc(report['generated_at'])} &nbsp;|&nbsp; "
        f"Target role: {esc(report['role'])} ({esc(report['seniority'])})",
        styles["Meta"]))
    flow.append(Paragraph(
        f"Industry / Company type / Company: {esc(report['industry'])} / "
        f"{esc(report['company_type'])} / {esc(report['company'])}",
        styles["Meta"]))
    flow.append(Spacer(1, 6))
    flow.append(HRFlowable(width="100%", color=colors.HexColor("#1E88E5")))
    flow.append(Spacer(1, 6))

    flow.append(Paragraph("Overall Result", styles["H2c"]))
    flow.append(Paragraph(
        f"<b>Overall interview score:</b> {report['overall_score']}/10<br/>"
        f"<b>Recommendation:</b> {esc(report['recommendation'])}<br/>"
        f"<b>CV–Role match:</b> {report['match_score']}%", styles["Body"]))

    # Narrative: convert simple Markdown to paragraphs.
    for block in _markdown_blocks(report["narrative_md"]):
        kind, text = block
        if kind == "h":
            flow.append(Paragraph(esc(text), styles["H2c"]))
        elif kind == "li":
            flow.append(Paragraph("• " + esc(text), styles["Body"]))
        else:
            flow.append(Paragraph(esc(text), styles["Body"]))

    flow.append(Spacer(1, 6))
    flow.append(Paragraph("Question-by-Question Breakdown", styles["H2c"]))
    for p in report["per_question"]:
        crit = ", ".join(f"{k.replace('_',' ')}: {v}/10" for k, v in p.get("criteria", {}).items())
        flow.append(Paragraph(
            f"<b>Q{p['index']} [{esc(p['category'])}] — {p['score']}/10</b>", styles["Body"]))
        flow.append(Paragraph(f"<b>Question:</b> {esc(p['question'])}", styles["Body"]))
        flow.append(Paragraph(f"<b>Answer:</b> {esc(p['answer'])}", styles["Body"]))
        if crit:
            flow.append(Paragraph(f"<b>Criteria:</b> {esc(crit)}", styles["Body"]))
        flow.append(Paragraph(f"<b>Strengths:</b> {esc('; '.join(p['strengths']) or '—')}", styles["Body"]))
        flow.append(Paragraph(f"<b>Weaknesses:</b> {esc('; '.join(p['weaknesses']) or '—')}", styles["Body"]))
        flow.append(Paragraph(f"<b>How to improve:</b> {esc('; '.join(p['improvements']) or '—')}", styles["Body"]))
        flow.append(Paragraph(f"<b>Feedback:</b> {esc(p['feedback'])}", styles["Body"]))
        flow.append(Spacer(1, 6))

    doc.build(flow)
    buffer.seek(0)
    return buffer.getvalue()


def build_docx(report):
    """Return the report as DOCX bytes using python-docx."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    document = Document()
    title = document.add_heading(f"Interview Report — {report['candidate']}", level=0)

    meta = document.add_paragraph()
    meta.add_run(
        f"Date: {report['generated_at']} | Target role: {report['role']} "
        f"({report['seniority']})\n").italic = True
    meta.add_run(
        f"Industry / Company type / Company: {report['industry']} / "
        f"{report['company_type']} / {report['company']}").italic = True

    document.add_heading("Overall Result", level=1)
    p = document.add_paragraph()
    p.add_run("Overall interview score: ").bold = True
    p.add_run(f"{report['overall_score']}/10\n")
    p.add_run("Recommendation: ").bold = True
    p.add_run(f"{report['recommendation']}\n")
    p.add_run("CV–Role match: ").bold = True
    p.add_run(f"{report['match_score']}%")

    for kind, text in _markdown_blocks(report["narrative_md"]):
        if kind == "h":
            document.add_heading(text, level=1)
        elif kind == "li":
            document.add_paragraph(text, style="List Bullet")
        else:
            document.add_paragraph(text)

    document.add_heading("Question-by-Question Breakdown", level=1)
    for p in report["per_question"]:
        document.add_heading(
            f"Q{p['index']} [{p['category']}] — {p['score']}/10", level=2)
        for label, value in [
            ("Question", p["question"]),
            ("Answer", p["answer"]),
            ("Strengths", "; ".join(p["strengths"]) or "—"),
            ("Weaknesses", "; ".join(p["weaknesses"]) or "—"),
            ("How to improve", "; ".join(p["improvements"]) or "—"),
            ("Feedback", p["feedback"]),
        ]:
            para = document.add_paragraph()
            para.add_run(f"{label}: ").bold = True
            para.add_run(str(value))

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _markdown_blocks(md):
    """Very small Markdown splitter -> list of (kind, text) tuples."""
    blocks = []
    for raw in (md or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("#"):
            blocks.append(("h", line.lstrip("#").strip()))
        elif line.startswith(("- ", "* ", "• ")):
            blocks.append(("li", line[2:].strip().lstrip("•").strip()))
        else:
            # strip bold markers for clean rendering
            blocks.append(("p", line.replace("**", "")))
    return blocks


def filename(report, ext):
    return f"{_safe_filename(report['candidate'])}.{ext}"
